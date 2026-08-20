# core/excel_exporter.py
from typing import List, Dict, Any
from core.persian_utils import format_toman, to_persian_numbers


def export_transactions_to_excel(transactions_data: List[Dict[str, Any]], output_path: str) -> bool:
    """تولید فایل خروجی اکسل از تراکنش‌ها با پشتیبانی از باز شدن در Excel"""
    try:
        # ۱. تلاش برای خروجی با openpyxl در صورت نصب بودن
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "دفتر مالی و تراکنش‌ها"
            ws.sheet_view.rightToLeft = True

            headers = ["کد تراکنش", "پروژه / عروس و داماد", "استودیو / آتلیه", "مبلغ واریزی (تومان)", "تاریخ واریز"]
            ws.append(headers)

            # استایل هدر
            header_fill = PatternFill(start_color="1E1E24", end_color="1E1E24", fill_type="solid")
            header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
            border_thin = Border(
                left=Side(style="thin", color="CCCCCC"),
                right=Side(style="thin", color="CCCCCC"),
                top=Side(style="thin", color="CCCCCC"),
                bottom=Side(style="thin", color="CCCCCC")
            )

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border_thin

            for row_idx, item in enumerate(transactions_data, start=2):
                ws.append([
                    f"TX-{item.get('id', 0):04d}",
                    item.get("project", "-"),
                    item.get("client", "-"),
                    item.get("amount", 0),
                    item.get("date", "-")
                ])
                for col_idx in range(1, len(headers) + 1):
                    c = ws.cell(row=row_idx, column=col_idx)
                    c.font = Font(name="Tahoma", size=10)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border = border_thin

            # تنظیم عرض ستون‌ها
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

            wb.save(output_path)
            return True

        except ImportError:
            # ۲. خروجی جایگزین CSV استاندارد سازگار با اکسل در صورت نبود openpyxl
            import csv
            with open(output_path.replace(".xlsx", ".csv"), "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["کد تراکنش", "پروژه / عروس و داماد", "استودیو / آتلیه", "مبلغ واریزی (تومان)", "تاریخ"])
                for item in transactions_data:
                    writer.writerow([
                        f"TX-{item.get('id', 0):04d}",
                        item.get("project", "-"),
                        item.get("client", "-"),
                        item.get("amount", 0),
                        item.get("date", "-")
                    ])
            return True

    except Exception as e:
        print(f"Excel Export Error: {e}")
        return False